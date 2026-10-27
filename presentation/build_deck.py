"""Build the project presentation (presentation/fuel_system_digital_twin.pptx).

Run from the repository root after the notebooks have produced results/figures:

    .venv\\Scripts\\python presentation\\build_deck.py
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
OUT = ROOT / "presentation" / "fuel_system_digital_twin.pptx"
REPO = "github.com/PeterImoniose/fuel-system-digital-twin"

NAVY = RGBColor(0x13, 0x23, 0x3F)
INK = RGBColor(0x1B, 0x1F, 0x24)
MUTED = RGBColor(0x5F, 0x66, 0x70)
LIGHT = RGBColor(0xFC, 0xFC, 0xFB)  # same surface as the notebook figures
CARD = RGBColor(0xEE, 0xF1, 0xF4)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREY = RGBColor(0x9A, 0xA3, 0xAD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xC9, 0xD6, 0xEA)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5
MARGIN = 0.6


# --- helpers ---------------------------------------------------------------------------
def new_slide(dark: bool = False):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY if dark else LIGHT
    return s


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    """Add a text box. ``runs`` is a string or a list of paragraphs (str or list of (text, opts))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        pieces = para if isinstance(para, list) else [(para, {})]
        for t, opt in pieces:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", italic)
            f.color.rgb = opt.get("color", color)
        if i < len(paras) - 1:
            p.space_after = Pt(opt.get("space", 6) if pieces else 6)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pieces = item if isinstance(item, list) else [(item, {})]
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.size = Pt(size - 4)
        r0.font.color.rgb = AMBER
        r0.font.name = FONT
        for t, opt in pieces:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", False)
            r.font.color.rgb = opt.get("color", color)
    return tb


def title(slide, t, dark=False, sub=None):
    text(slide, MARGIN, 0.45, W - 2 * MARGIN, 0.8, t, size=34, bold=True, color=WHITE if dark else INK,
         name="Title")
    if sub:
        text(slide, MARGIN, 1.2, W - 2 * MARGIN, 0.45, sub, size=16, color=ICE if dark else MUTED, name="Subtitle")


def card(slide, x, y, w, h, fill=CARD):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def stat(slide, x, y, w, big, label, dark=False, big_color=None, big_size=40, h=1.55):
    card(slide, x, y, w, h, fill=RGBColor(0x1E, 0x33, 0x58) if dark else CARD)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.75, big, size=big_size, bold=True,
         color=big_color or (AMBER if dark else NAVY))
    text(slide, x + 0.25, y + 0.92, w - 0.5, h - 1.0, label, size=13, color=ICE if dark else MUTED)


def badge(slide, x, y, label, d=0.55, fill=AMBER, color=NAVY, size=16):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color


def image(slide, path, x, y, w, h, align="center"):
    """Place an image inside the box, keeping its aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    dx = x + (w - dw) / 2 if align == "center" else x
    dy = y + (h - dh) / 2
    pic = slide.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
    return pic


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def footer(slide, n, dark=False):
    text(slide, W - MARGIN - 1.0, H - 0.45, 1.0, 0.3, str(n), size=11, color=ICE if dark else GREY,
         align=PP_ALIGN.RIGHT)


def style_chart(chart, size=12):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = MUTED


# --- 1. title ------------------------------------------------------------------------------
s = new_slide(dark=True)
badge(s, MARGIN, 1.55, "", d=0.35)
text(s, MARGIN, 2.1, 11.5, 1.2, "Fuel-System Digital Twin", size=54, bold=True, color=WHITE)
text(s, MARGIN, 3.25, 11.5, 0.9, "Detecting and diagnosing faults in an aircraft fuel rig with physics, "
     "machine learning and honest validation", size=22, color=ICE)
text(s, MARGIN, 5.6, 8, 0.4, "Avwerosuo Peter Imoniose  |  Data science project", size=16, color=WHITE)
text(s, MARGIN, 6.05, 8, 0.4, REPO, size=14, color=AMBER)
notes(s, "A data-science project on a laboratory fuel-system rig: can we tell from a few sensors whether "
         "it is healthy, which component is failing, how badly, and how early, and can we trust the answer?")

# --- 2. the problem ---------------------------------------------------------------------------
s = new_slide()
title(s, "Five ways a fuel system fails, one set of sensors",
      sub="A gear pump feeds a filter, a shut-off valve and an injector; five valves emulate the faults")
image(s, FIG / "01_rig_schematic.png", MARGIN, 1.85, 12.1, 3.4)
qs = [("1", "Is the system healthy?"), ("2", "Which component is failing, and how badly?"),
      ("3", "How early, and with what confidence?")]
for i, (n, q) in enumerate(qs):
    x = MARGIN + i * 4.1
    badge(s, x, 5.75, n)
    text(s, x + 0.75, 5.75, 3.2, 0.6, q, size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)
notes(s, "Clogged filter, degraded pump, stuck valve, leaking pipe and clogged injector, each emulated by a "
         "proportional valve at increasing severity. Sensors: five pressures, one flow meter at the end of "
         "the line, and the pump speed.")

# --- 3. the data ------------------------------------------------------------------------------------
s = new_slide()
title(s, "The data: 2.8 million sensor readings and 22 pump tests")
stats = [("47", "rig recordings, 60 s each at 1 kHz: 6 healthy, 41 faulty"),
         ("5", "fault types at up to 9 severities each (10-100 % valve travel)"),
         ("2.8 M", "samples, averaged into 2 820 one-second windows"),
         ("22", "pump tests at 400-1140 rpm from two campaigns (2021, 2023)")]
for i, (b, l) in enumerate(stats):
    stat(s, MARGIN + i * 3.06, 1.75, 2.85, b, l, h=1.9)
bullets(s, MARGIN, 4.2, 12, 2.6, [
    [("Data audit first: ", {"bold": True}), ("mixed gauge and absolute pressures, a sensor clipped at zero, "
                                             "mislabelled columns and a header that hides the real start time.", {})],
    [("Python package ", {"bold": True}), ("for loading and cleaning (LabVIEW files, units, derived pressure drops), "
                                          "with unit tests.", {})],
    [("Eight notebooks ", {"bold": True}), ("from audit to live monitoring, all rebuilt from raw data with one command.", {})],
], size=17)
footer(s, 3)
notes(s, "The audit turned up issues that would silently corrupt an analysis: pressures on two different "
         "references, a clipped sensor, mislabelled channels. These became data-quality checks in the final system.")

# --- 4. the catch -----------------------------------------------------------------------------------
s = new_slide()
title(s, "The catch: every fault was recorded at a different time of day")
image(s, FIG / "01_recording_timeline.png", MARGIN, 1.55, 8.2, 3.4)
stat(s, 9.2, 1.7, 3.5, "96 %", "of windows named correctly by a model that sees only the clock", big_size=48, h=1.9)
bullets(s, MARGIN, 5.25, 12.1, 1.9, [
    "Healthy data exists only for 09:58-10:03; faults follow in blocks until 16:11.",
    [("So accuracy alone proves nothing here: ", {"bold": True}),
     ("the key results come from tests the clock cannot help with.", {})],
], size=17)
footer(s, 4)
notes(s, "A random forest given only the time of day scores 96 % when each recording is held out, better than any "
         "sensor model. This is the single most important finding about the data, and it shaped every validation choice.")

# --- 5. healthy baseline --------------------------------------------------------------------------
s = new_slide()
title(s, "Normal is repeatable, and the noise is the pump itself")
image(s, FIG / "02_pressure_profile.png", MARGIN, 1.5, 7.4, 4.3)
stat(s, 8.45, 1.55, 4.3, "0.06 %", "variation in flow between healthy runs", h=1.45)
stat(s, 8.45, 3.15, 4.3, "~75 %", "of pressure fluctuation sits at the pump's shaft frequency and harmonics", h=1.45)
stat(s, 8.45, 4.75, 4.3, "1-3.5 %", "drop in pump delivery by the afternoon: drift, not a fault", h=1.45)
text(s, MARGIN, 6.45, 12, 0.6, "Healthy limits belong on 1 s averages: limits on raw samples are 2-14 times wider and hide "
     "small faults.", size=16, color=MUTED, italic=True)
footer(s, 5)
notes(s, "Pressure falls along the line through the filter, rises 0.47 bar across the pump, then falls across the "
         "valve and nozzle. Frequency analysis shows the noise is pump ripple at 6.7 Hz and a gear-mesh line at 80 Hz. "
         "The healthy point also drifts through the day, about as much as the mildest faults.")

# --- 6. physics indicators ----------------------------------------------------------------------------
s = new_slide()
title(s, "Physics-based indicators tell all five faults apart")
image(s, FIG / "03_residual_heatmap.png", 9.0, 1.35, 3.8, 5.7)
bullets(s, MARGIN, 1.5, 7.6, 3.4, [
    [("One indicator per pipe section: ", {"bold": True}),
     ("measured pressure drop against what a healthy section would show at that flow.", {})],
    [("Fitted without the fault being scored ", {"bold": True}),
     ("(leave-one-fault-out), so no fault ever sees its own data.", {})],
    [("Each fault lights up its own pattern, ", {"bold": True}),
     ("including filter, pump and leak faults that the rig's expert rules cannot separate.", {})],
], size=17)
card(s, MARGIN, 5.55, 7.6, 1.4)
text(s, MARGIN + 0.3, 5.68, 7.0, 1.2, [
    [("Key insight: ", {"bold": True, "color": NAVY}),
     ("leak flow still passes through the shut-off valve before escaping, so the valve's pressure drop rises for a "
      "leak but falls for a pump fault. One physical fact resolves the ambiguity.", {})]], size=16)
footer(s, 6)
notes(s, "Each section obeys a pressure-flow law while healthy. Residuals against those laws, plus a pump residual for "
         "flow that never reaches the meter, give every fault a distinct pattern. The heatmap shows all 47 recordings.")

# --- 7. degradation paths ---------------------------------------------------------------------------
s = new_slide()
title(s, "Blockages slide along the pump curve; lost flow falls below it")
image(s, FIG / "03_degradation_paths.png", MARGIN, 1.4, 8.3, 5.4)
bullets(s, 9.2, 1.7, 3.6, 5.0, [
    [("Filter, valve, nozzle: ", {"bold": True}), ("higher pressure, lower flow, on the pump's own curve.", {})],
    [("Pump bypass and leak: ", {"bold": True}), ("lower pressure and lower flow: part of the flow never arrives.", {})],
    [("The 2014 pump curve ", {"bold": True}), ("is steeper than in 2021-23, so a pump model must be recalibrated "
                                               "on the asset it monitors.", {})],
], size=16, gap=14)
footer(s, 7)
notes(s, "Plotting every recording on the pump's pressure-flow plane shows each fault as a path. This view explains "
         "why restriction faults and lost-flow faults look different to the pump.")

# --- 8. diagnosis -----------------------------------------------------------------------------------
s = new_slide()
title(s, "Fault diagnosis: the split matters more than the model",
      sub="Macro F1 on 1 s windows. Grey: leave one recording out. Amber: train on severe faults, test on mild ones")
cd = CategoryChartData()
methods = ["Expert rule tree", "Logistic regression (raw)", "Signature direction", "Physics rules (never saw the fault)",
           "Gradient boosting (physics features)", "Clock time only (shortcut)"]
cd.categories = methods
cd.add_series("Leave one recording out", (0.471, 0.873, 0.868, 0.880, 0.921, 0.958))
cd.add_series("Train severe, test mild", (0.250, 0.056, 0.576, 0.580, 0.547, 0.603))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(MARGIN), Inches(1.8), Inches(8.2), Inches(5.2), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
for ser, col in zip(ch.series, (GREY, AMBER)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
    ser.data_labels.show_value = True
    ser.data_labels.number_format = "0.00"
    ser.data_labels.number_format_is_linked = False
    ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser.data_labels.font.size = Pt(11)
    ser.data_labels.font.color.rgb = INK
ch.plots[0].gap_width = 60
ch.plots[0].overlap = -10
va = ch.value_axis
va.maximum_scale, va.minimum_scale = 1.1, 0
va.has_major_gridlines = False
va.visible = False
ch.category_axis.tick_labels.font.size = Pt(12)
ch.category_axis.format.line.fill.background()
bullets(s, 9.2, 1.9, 3.6, 5.0, [
    [("Shuffled splits score 1.00 ", {"bold": True}), ("for every tree model: windows of one recording are near copies.", {})],
    [("Standard classifiers collapse ", {"bold": True}), ("on mild faults: they learn distance from healthy, not direction.", {})],
    [("Physics rules hold up ", {"bold": True}), ("with no clock and no data from the fault: the cleanest evidence.", {})],
], size=15, gap=14)
footer(s, 8)
notes(s, "Leave-one-recording-out scores are high for many models, but the clock alone does as well. Training on clear "
         "faults and testing on mild ones separates the approaches: methods that use the direction of a deviation "
         "generalise, standard classifiers do not.")

# --- 9. sim-to-real -----------------------------------------------------------------------------------
s = new_slide()
title(s, "Simulated faults can stand in for real fault data",
      sub="A classifier trained only on faults generated by a calibrated model of the rig, tested on the real recordings")
stat(s, MARGIN, 1.85, 3.7, "0.87", "macro F1 on real data (+/- 0.03 over five seeds)", big_size=46, h=1.6)
stat(s, MARGIN, 3.6, 3.7, "40-42 / 47", "real recordings named correctly", big_size=34, h=1.6)
stat(s, MARGIN, 5.35, 3.7, "0", "real fault labels used in training", big_size=40, h=1.45)
image(s, FIG / "06_sim_to_real_confusion.png", 4.7, 1.8, 4.6, 4.95)
bullets(s, 9.6, 1.95, 3.2, 4.9, [
    [("Misses are the mildest settings, ", {"bold": True}), ("called healthy: the safe direction.", {})],
    [("Caveat: ", {"bold": True}), ("the healthy boundary does not transfer yet (some healthy windows flagged, "
                                    "depending on the seed).", {})],
    [("So it runs as a second opinion, ", {"bold": True}), ("never raising an alarm on its own.", {})],
], size=15, gap=14)
footer(s, 9)
notes(s, "A steady-state network model joins the identified component laws. Faults are injected over wide ranges with "
         "random drift and noise, and a gradient-boosting classifier is trained on the synthetic data only.")

# --- 10. pump --------------------------------------------------------------------------------------
s = new_slide()
title(s, "The pump matches its datasheet; a relief valve caps pressure")
image(s, FIG / "05_pump_curves.png", MARGIN, 1.4, 8.4, 3.45)
stat(s, MARGIN, 5.15, 2.7, "< 2.3 %", "from the OEM datasheet at 1140 rpm", h=1.6)
stat(s, MARGIN + 2.85, 5.15, 2.7, "3.45 bar", "built-in relief valve: flow collapses past it", h=1.6)
stat(s, MARGIN + 5.7, 5.15, 2.7, "~1 %", "physics model error on an unseen speed", h=1.6)
bullets(s, 9.3, 1.6, 3.5, 5.2, [
    [("2021 tests ", {"bold": True}), ("closed the valves to 10 % open and pushed past the relief valve; 2023 tests "
                                     "stopped at 25-35 %.", {})],
    [("Displacement is constant ", {"bold": True}), ("(2.1 mL/rev); internal leakage changes from day to day.", {})],
    [("Day-to-day shifts of 3-10 % ", {"bold": True}), ("dominate real error, so the model keeps displacement fixed "
                                                       "and re-estimates leakage.", {})],
    [("Found an error ", {"bold": True}), ("in the published verification figures: 3.6 %, not under 1.5 %.", {})],
], size=15, gap=12)
footer(s, 10)
notes(s, "Straight-line fits per speed give displacement and slip. The pump matches its datasheet; the flow collapse "
         "in the first campaign is the relief valve opening at 50 psig.")

# --- 11. regression ----------------------------------------------------------------------------------
s = new_slide()
title(s, "Regression: test on what the model has not seen",
      sub="Pump flow from speed and pressure; error when a whole pump speed is held out (L/min, lower is better)")
cd = CategoryChartData()
cats = ["Random forest", "Gradient boosting", "Neural network", "Gaussian process", "Physics (2 parameters)",
        "Physics + Gaussian process", "Polynomial, degree 2", "Polynomial, degree 3"]
vals = (0.1124, 0.1097, 0.0632, 0.0172, 0.0156, 0.0135, 0.0130, 0.0105)
cd.categories = cats
cd.add_series("RMSE on a held-out speed", vals)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(MARGIN), Inches(1.8), Inches(6.6), Inches(5.0), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = False
ser = ch.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = BLUE
for i, v in enumerate(vals):
    if v > 0.05:
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = AMBER
ser.data_labels.show_value = True
ser.data_labels.number_format = "0.000"
ser.data_labels.number_format_is_linked = False
ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser.data_labels.font.size = Pt(11)
ser.data_labels.font.color.rgb = INK
ch.plots[0].gap_width = 50
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 0.135
ch.category_axis.tick_labels.font.size = Pt(12)
ch.category_axis.format.line.fill.background()
image(s, FIG / "07_regressors_extrapolation.png", 7.5, 1.85, 5.25, 3.0)
bullets(s, 7.5, 5.0, 5.25, 2.0, [
    [("Trees cannot predict an unseen speed: ", {"bold": True}), ("they copy the nearest one.", {})],
    [("Training R² kept rising with polynomial degree ", {"bold": True}), ("while accuracy on new conditions fell.", {})],
    [("The hybrid adds prediction intervals, ", {"bold": True}), ("which hold within a test day but not across days.", {})],
], size=14, gap=8)
footer(s, 11)
notes(s, "The earlier version of this project reported R-squared 0.997 on a shuffled split. Testing on held-out speeds, "
         "independent points and another campaign shows which models really generalise.")

# --- 12. monitoring -----------------------------------------------------------------------------------
s = new_slide()
title(s, "Live monitoring: catch the change, not the anomaly",
      sub="Calibrated on healthy data only, then run on streams of real recordings")
image(s, FIG / "08_change_detection_example.png", MARGIN, 1.75, 7.6, 3.45)
stat(s, 8.6, 1.8, 4.15, "< 1 s", "to catch each real step of degradation (CUSUM and Bayesian change points)", h=1.55)
stat(s, 8.6, 3.5, 4.15, "1-3 %", "extra resistance when a slow, simulated clog is first flagged", h=1.55)
bullets(s, MARGIN, 5.5, 12.1, 1.6, [
    [("Anomaly detectors trained on 5 minutes of healthy data flag everything that follows ", {"bold": True}),
     ("(even settings with no fault): a data-coverage limit, not an algorithm one.", {})],
    [("Kalman-smoothed indicators ", {"bold": True}), ("lift accuracy on borderline recordings by 10-15 points.", {})],
], size=15, gap=8)
footer(s, 12)
notes(s, "Comparing against the recent past works where comparing against a fixed healthy reference fails. The Bayesian "
         "detector raised no false alarms on healthy streams.")

# --- 13. digital shadow --------------------------------------------------------------------------------
s = new_slide()
title(s, "Putting it together: the whole test day replayed live")
image(s, FIG / "06_day_replay.png", MARGIN, 1.4, 12.1, 3.75)
stat(s, MARGIN, 5.35, 3.85, "0", "false alarms on healthy data", h=1.5)
stat(s, MARGIN + 4.12, 5.35, 3.85, "41 / 47", "recordings end on the right condition", h=1.5)
stat(s, MARGIN + 8.25, 5.35, 3.85, "100 %", "of the one drift-caused mistake flagged as low confidence", h=1.5)
footer(s, 13)
notes(s, "The digital shadow combines the physics indicators, the rules, a simulation-trained second opinion, a "
         "5-second persistence rule and a green/yellow/red quality indicator. The misses are settings with no physical "
         "effect, the two mildest blockages, and one case of drift.")

# --- 14. next steps -----------------------------------------------------------------------------------
s = new_slide(dark=True)
title(s, "What limits it is the data, not the algorithms", dark=True,
      sub="Next steps that would turn the digital shadow into a digital twin")
steps = [("Measure fluid temperature", "It drives the pump drift that mimics faults."),
         ("Record healthy data all day", "Five minutes cannot define normal."),
         ("Randomise the fault schedule", "Removes the time-of-day shortcut."),
         ("Fix the sensor gaps", "Log air pressure; widen the clipped sensor's range."),
         ("Add a second flow meter", "Measures lost flow directly at the pump."),
         ("Close the loop", "Feed actions back, such as pump speed or isolating a leak.")]
for i, (h_, d_) in enumerate(steps):
    col, row = i % 3, i // 3
    x, y = MARGIN + col * 4.1, 2.05 + row * 2.45
    card(s, x, y, 3.85, 2.15, fill=RGBColor(0x1E, 0x33, 0x58))
    badge(s, x + 0.3, y + 0.3, str(i + 1), d=0.5, size=15)
    text(s, x + 0.95, y + 0.3, 2.75, 0.55, h_, size=17, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, y + 1.05, 3.3, 1.0, d_, size=14, color=ICE)
footer(s, 14, dark=True)
notes(s, "Every limit met in the project points to a concrete fix, mostly in what data is collected rather than in "
         "the models.")

# --- 15. close ------------------------------------------------------------------------------------------
s = new_slide(dark=True)
text(s, MARGIN, 1.0, 12, 0.9, "Tools and skills", size=34, bold=True, color=WHITE)
skills = [("Data engineering", "Python package, LabVIEW parsing, unit handling, data audit, pytest"),
          ("Analysis", "pandas, NumPy, SciPy, spectral analysis, physics-based feature design"),
          ("Machine learning", "scikit-learn classifiers and regressors, Gaussian processes, anomaly detection"),
          ("Validation", "leave-one-group-out, shortcut tests, train-severe/test-mild, independent test sets"),
          ("Monitoring", "CUSUM, Bayesian change points, Kalman filters, simulation-based training"),
          ("Delivery", "executed notebooks, Streamlit app, reproducible build, this deck")]
for i, (a, b) in enumerate(skills):
    col, row = i % 2, i // 2
    x, y = MARGIN + col * 6.1, 2.1 + row * 1.25
    badge(s, x, y + 0.05, "", d=0.32)
    text(s, x + 0.55, y, 5.3, 0.4, a, size=18, bold=True, color=WHITE)
    text(s, x + 0.55, y + 0.42, 5.3, 0.7, b, size=14, color=ICE)
text(s, MARGIN, 6.25, 12, 0.45, [[("Code, notebooks and app: ", {"color": WHITE}), (REPO, {"color": AMBER, "bold": True})]],
     size=18)
notes(s, "Everything is reproducible from the raw data with one command; the repository has the notebooks, the "
         "package, tests and an interactive app.")

prs.save(OUT)
print(f"saved {OUT} with {len(prs.slides)} slides")
